#!/usr/bin/env python3
# includes/import_excel.py
# Použití: python3 import_excel.py preview <file.xlsx|file.csv>
#          python3 import_excel.py import  <file.xlsx|file.csv>
import sys, json, math, csv
from datetime import datetime

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

try:
    import openpyxl # type: ignore
except ImportError:
    openpyxl = None

def safe(val):
    if val is None: return None
    if isinstance(val, float) and math.isnan(val): return None
    if isinstance(val, datetime): return val.strftime('%Y-%m-%d')
    return str(val).strip() if str(val).strip() not in ('None','nan','') else None

def safe_int(val, default=0):
    try:
        v = float(val) if val is not None else 0
        return int(v) if not math.isnan(v) else default
    except: return default

def safe_float(val):
    try:
        v = float(val)
        return None if math.isnan(v) else round(v, 8)
    except: return None

def safe_date(val):
    if val is None: return None
    if isinstance(val, datetime): return val.strftime('%Y-%m-%d')
    try:
        return datetime.strptime(str(val)[:10], '%Y-%m-%d').strftime('%Y-%m-%d')
    except: return None

def normalize_smena(val):
    if not val: return 'O8'
    v = str(val).strip().upper()
    for s in ['R8','O8','N8','R12','N12']:
        if s in v: return s
    return 'O8'

def parse_row(hdr, row):
    """Map Excel columns to DB fields"""
    def col(name):
        try:
            idx = hdr.index(name)
            return row[idx] if idx < len(row) else None
        except ValueError: return None

    # Column aliases from DVD-2025
    datum  = safe_date(col('Datum'))
    if not datum: return None

    smena_raw = col('Směna \n(R-O-N)') or col('Směna') or ''
    vyrobeno  = safe_int(col('Vyrobeno [ks]'))
    nok       = safe_int(col('NOK kusů'))
    dobre     = safe_int(col('Vyrobeno dobře \n[ks]')) or max(0, vyrobeno - nok)

    oee_a = safe_float(col('Avialability\nDostupnost pracoviště'))
    oee_p = safe_float(col('Performance (Výkon pracoviště)'))
    oee_q = safe_float(col('Kvalita'))
    oee   = safe_float(col('OEE %'))

    # prostoje
    def pm(k): return safe_int(col(k))

    p1  = pm('P1\n změna výroby - verze, formy')
    p2  = pm('P2 rozjezd linky')
    p3  = pm('P3\nVzorkování, testování,DUMMY')
    p4  = pm('P4\nčištění, údržba')
    p5  = pm('P5 schůzky, školení')
    p6  = pm('P6 školitel')
    p7  = pm('P7 zaškolovaný')
    p8  = pm('P8 změna pracoviště')
    p9  = pm('P9 Výpomoc, inventura')
    p10 = pm('P10\nplánovaná odstávka')
    o1  = pm('O1 lékař, nevolnost ')
    l1  = pm('L1  chybějící materiál sklad')
    l2  = pm('L2  čekání na materiál (manipulant, skladník)')
    v1  = pm('V1 organizační problémy  chybí pracovník')
    v2  = pm('V2\n chybějící materiál předchozí pracoviště')
    v3  = pm('V3\n seřízení linky')
    q1  = pm('Q1 \n Nekvalitní vstupní materiál')
    q2  = pm('Q2 \n Změna šarže   (požadavek kvality)')
    q3  = pm('Q3 \n sorting')
    q4  = pm('Q4\n uvolnění výroby')

    pt  = safe_int(col('Doba prostoje [min]\nTechnicky'))
    pv  = safe_int(col('Doba prostoje [min]\nVýrobní'))
    pk  = safe_int(col('Doba prostoje [min]\nKvalitativní'))
    pl  = safe_int(col('Doba prostoje [min]\nLogistická'))
    po  = safe_int(col('Doba prostoje [min]\nOstatní'))
    pp  = safe_int(col('Doba prostoje [min]\nplánované'))
    pnp = safe_int(col('Doba prostoje [min]\nneplánované'))
    pc  = safe_int(col('Doba prostoje [min]\ncelkem'))

    norma = safe_float(col('Norma'))
    cycle = safe_float(col('Cycle time/ díl (min.)'))
    final = safe(col('Final'))

    try: cw  = int(col('CW') or 0)
    except: cw = 0
    try: rok = int(col('rok') or datetime.strptime(datum,'%Y-%m-%d').year)
    except: rok = 0
    try: mes = int(col('měsíc') or datetime.strptime(datum,'%Y-%m-%d').month)
    except: mes = 0

    cisty = safe_int(col('Čistý čas [min]')) or safe_int(col('pracovní doba'))
    prac  = safe_int(col('pracovní doba')) or 480
    prest = safe_int(col('Přesčas'))

    zakaznik = safe(col('Zákazník')) or safe(col('Zákarník'))
    jmeno    = safe(col('Jméno'))
    zapsal   = safe(col('Zapsal'))

    # typ detection
    pracovisite_raw = safe(col('Pracoviště')) or ''
    nazev_dilu = safe(col('Název dílu')) or ''
    typ = 'filtry' if any(x in nazev_dilu.lower() for x in ['filter','filtr','emc']) else 'serie'

    return {
        ':typ':   typ,
        ':datum': datum, ':cw': cw, ':rok': rok, ':mesic': mes,
        ':smena': normalize_smena(smena_raw),
        ':zakaznik': zakaznik, ':projekt': safe(col('Projekt')),
        ':verze': safe(col('Verze')),
        ':kp':   pracovisite_raw,
        ':np':   safe(col('Název pracoviště')),
        ':nd':   nazev_dilu,
        ':sap':  safe(col('SAP')),
        ':ns':   safe_int(col('Náklanové středisko')) or None,
        ':jmeno': jmeno, ':pozice': safe(col('Pracovní pozice')),
        ':vyr': vyrobeno, ':nok': nok, ':pos': safe_int(col('K posouzení')),
        ':dobre': dobre,
        ':cisty': cisty, ':prac_doba': prac, ':prestas': prest,
        ':p1':p1,':p2':p2,':p3':p3,':p4':p4,':p5':p5,
        ':p6':p6,':p7':p7,':p8':p8,':p9':p9,':p10':p10,
        ':o1':o1,':l1':l1,':l2':l2,
        ':v1':v1,':v2':v2,':v3':v3,
        ':q1':q1,':q2':q2,':q3':q3,':q4':q4,
        ':pt':pt,':pv':pv,':pk':pk,':pl':pl,':po':po,
        ':pp':pp,':pnp':pnp,':pc':pc,
        ':oee_a': oee_a, ':oee_p': oee_p, ':oee_q': oee_q, ':oee': oee,
        ':norma': norma, ':cycle': cycle, ':final': final,
        ':zapsal': zapsal,
        # For preview
        'datum': datum, 'zakaznik': zakaznik or '',
        'nazev_dilu': nazev_dilu, 'smena': normalize_smena(smena_raw),
        'vyr': vyrobeno, 'nok': nok, 'oee': oee,
    }

def load_wb(path):
    if openpyxl is None:
        raise RuntimeError("openpyxl není nainstalován: pip install openpyxl")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # Try sheet named 'Databáze' first, else first sheet
    ws = None
    for name in wb.sheetnames:
        if 'Databáze' in name or 'Databaze' in name.replace('á','a'):
            ws = wb[name]; break
    if ws is None: ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows: raise ValueError("Prázdný list")
    hdr = [str(c).strip() if c else '' for c in rows[0]]
    return hdr, rows[1:]

def load_csv(path):
    encodings = ['utf-8-sig', 'cp1250', 'iso-8859-2']
    last_error = None

    for enc in encodings:
        try:
            with open(path, 'r', encoding=enc, newline='') as fh:
                sample = fh.read(4096)
                fh.seek(0)
                delimiter = ';' if sample.count(';') >= sample.count(',') else ','
                reader = csv.reader(fh, delimiter=delimiter)
                rows = [row for row in reader if any(str(c).strip() for c in row)]

            if not rows:
                raise ValueError("Prázdný CSV soubor")

            hdr = [str(c).strip().lstrip('\ufeff') if c else '' for c in rows[0]]
            return hdr, rows[1:]
        except UnicodeDecodeError as e:
            last_error = e

    raise ValueError(f"CSV nelze dekódovat ({last_error})")

def load_source(path):
    p = path.lower()
    if p.endswith('.csv'):
        return load_csv(path)
    return load_wb(path)

def run_preview(path):
    hdr, data_rows = load_source(path)
    preview = []
    for row in data_rows[:20]:
        r = parse_row(hdr, row)
        if r: preview.append(r)
    print(json.dumps({'rows': preview, 'total_hdr': len(hdr)}, ensure_ascii=False))

def run_import(path):
    hdr, data_rows = load_source(path)
    rows = []
    skipped = 0
    for row in data_rows:
        r = parse_row(hdr, row)
        if r: rows.append(r)
        else: skipped += 1
    print(json.dumps({
        'imported': len(rows), 'skipped': skipped,
        'db_count': len(rows), 'rows': rows
    }, ensure_ascii=False, default=str))

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(json.dumps({'error': 'Použití: script.py preview|import <soubor.xlsx>'}))
        sys.exit(1)
    mode, fpath = sys.argv[1], sys.argv[2]
    try:
        if mode == 'preview': run_preview(fpath)
        elif mode == 'import': run_import(fpath)
        else: print(json.dumps({'error': 'Neznámý režim: '+mode}))
    except Exception as e:
        print(json.dumps({'error': str(e)}))
