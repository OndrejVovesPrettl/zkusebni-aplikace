import openpyxl
from pathlib import Path

wb = openpyxl.load_workbook(Path('uploads/prestavky.xlsx'), data_only=False)
ws = wb['Předání směn (2)']
print('SHEET:', ws.title, 'max', ws.max_row, ws.max_column)
for r in range(10, 50):
    items = []
    for c in range(1, 60):
        cell = ws.cell(r, c)
        v = cell.value
        if v is None or (isinstance(v, str) and v.strip()==''):
            continue
        items.append((c, repr(v), cell.number_format))
    if items:
        print('ROW', r, items)
