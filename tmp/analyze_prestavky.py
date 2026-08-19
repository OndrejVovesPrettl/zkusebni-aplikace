import openpyxl
from pathlib import Path

p = Path('uploads/prestavky.xlsx')
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb['Přestávky'] if 'Přestávky' in wb.sheetnames else wb.active

print('SHEET:', ws.title)

def val(cell):
    v = ws[cell].value
    if v is None:
        return ''
    return str(v).strip()

# Diagnostic rows from known table areas
for r in [13,14,15,17,18,19,20,25,26,27,30,31,32]:
    row_vals = []
    for c in range(1, 70):
        vv = ws.cell(r, c).value
        if vv is not None and str(vv).strip() != '':
            row_vals.append((c, str(vv).strip()))
    print(f'ROW {r}:', row_vals)

# count blocks by ranges based on template layout
# 1st floor 12h uses rows 18-20 (labels/times), 2nd floor 12h uses 30-32
for name, label_row, min_row, time_row in [
    ('zona1_12h', 18, 19, 20),
    ('zona2_12h', 30, 31, 32),
    ('zona1_8h', 13, 14, 15),
    ('zona2_8h', 25, 26, 27),
]:
    # each break pair is usually duration in row min_row and time in row time_row
    # count non-empty time cells in first 60 columns
    pairs = []
    for c in range(1, 60):
        t = ws.cell(time_row, c).value
        m = ws.cell(min_row, c).value
        if t is None:
            continue
        tt = str(t).strip()
        if not tt or '-' not in tt:
            continue
        mm = '' if m is None else str(m).strip()
        pairs.append((c, tt, mm))
    print(name, 'COUNT=', len(pairs))
    print(name, 'PAIRS=', pairs)
