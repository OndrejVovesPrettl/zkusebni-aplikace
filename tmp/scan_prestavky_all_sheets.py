import openpyxl, re
from pathlib import Path

wb = openpyxl.load_workbook(Path('uploads/prestavky.xlsx'), data_only=True)
for ws in wb.worksheets:
    print('\n===', ws.title, '===')
    time_pat = re.compile(r'\b\d{1,2}:\d{2}\s*-\s*\d{1,2}:\d{2}\b')
    found = []
    for r in range(1, min(ws.max_row, 120)+1):
        row_hits = []
        for c in range(1, min(ws.max_column, 80)+1):
            v = ws.cell(r,c).value
            if v is None:
                continue
            s = str(v).strip()
            if time_pat.search(s):
                row_hits.append((c, s))
        if row_hits:
            found.append((r, row_hits))
    for r, hits in found[:40]:
        print('ROW', r, hits)
