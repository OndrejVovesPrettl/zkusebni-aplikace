#!/usr/bin/env python3
# scripts/import_rework_excel.py
# Imports historical rework data from uploads/Opravy.xlsx into rework_entries
# Usage: python scripts/import_rework_excel.py [--dry-run]
import sys, os, math, re
from datetime import datetime

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

try:
    import mysql.connector
except ImportError:
    print("ERROR: mysql-connector-python not installed. Run: pip install mysql-connector-python")
    sys.exit(1)

# ── Config ─────────────────────────────────────────────────────────────────────
DB_HOST = 'localhost'
DB_NAME = 'dvd_vyroba'
DB_USER = 'root'
DB_PASS = 'InitAdmin123+'

EXCEL_PATH = os.path.join(os.path.dirname(__file__), '..', 'uploads', 'Opravy.xlsx')
VSTUP_SHEET = 'Vstup'

DRY_RUN = '--dry-run' in sys.argv

# ── Helpers ────────────────────────────────────────────────────────────────────
def safe(val):
    if val is None: return None
    if isinstance(val, float) and math.isnan(val): return None
    s = str(val).strip()
    return s if s not in ('None', 'nan', '') else None

def safe_date(val):
    if val is None: return None
    if isinstance(val, datetime): return val.strftime('%Y-%m-%d %H:%M:%S')
    try:
        s = str(val).strip()
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d.%m.%Y %H:%M:%S', '%d.%m.%Y'):
            try:
                return datetime.strptime(s[:len(fmt)], fmt).strftime('%Y-%m-%d %H:%M:%S')
            except ValueError:
                continue
    except Exception:
        pass
    return None

# ── Connect ────────────────────────────────────────────────────────────────────
print(f"Connecting to MySQL {DB_HOST}/{DB_NAME}...")
try:
    conn = mysql.connector.connect(
        host=DB_HOST, database=DB_NAME,
        user=DB_USER, password=DB_PASS,
        charset='utf8mb4'
    )
    cursor = conn.cursor(dictionary=True)
    print("Connected.")
except mysql.connector.Error as e:
    print(f"ERROR connecting to DB: {e}")
    sys.exit(1)

# ── Load codebooks from DB ─────────────────────────────────────────────────────
cursor.execute('SELECT id, code, label FROM rework_projects ORDER BY id')
projects_by_code = {}
for row in cursor.fetchall():
    projects_by_code[row['code'].strip().lower()] = int(row['id'])
    projects_by_code[row['label'].strip().lower()] = int(row['id'])
print(f"Projects: {list(projects_by_code.keys())}")

cursor.execute('SELECT id, code, label FROM rework_reasons ORDER BY id')
reasons = cursor.fetchall()
reasons_by_code = {}
for row in reasons:
    reasons_by_code[row['code'].strip().lower()] = int(row['id'])
    # Also index by code prefix (e.g., "01" from "01 - Description")
    code_num = re.match(r'^(\d+)', row['code'].strip())
    if code_num:
        reasons_by_code[code_num.group(1).lstrip('0') or '0'] = int(row['id'])
        reasons_by_code[code_num.group(1)] = int(row['id'])
print(f"Reasons loaded: {len(reasons)}")

cursor.execute('SELECT id, code, label FROM rework_actions ORDER BY id')
actions_by_code = {}
for row in cursor.fetchall():
    actions_by_code[row['code'].strip().lower()] = int(row['id'])
    actions_by_code[row['label'].strip().lower()] = int(row['id'])
print(f"Actions: {list(actions_by_code.keys())}")

# ── Open Excel ─────────────────────────────────────────────────────────────────
print(f"\nOpening Excel: {EXCEL_PATH}")
try:
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
except FileNotFoundError:
    print(f"ERROR: File not found: {EXCEL_PATH}")
    sys.exit(1)

if VSTUP_SHEET not in wb.sheetnames:
    print(f"ERROR: Sheet '{VSTUP_SHEET}' not found. Available: {wb.sheetnames}")
    sys.exit(1)

ws = wb[VSTUP_SHEET]
rows_iter = ws.iter_rows(values_only=True)

# Read header row
header_raw = next(rows_iter, None)
if header_raw is None:
    print("ERROR: Sheet is empty")
    sys.exit(1)

header = [str(h).strip() if h is not None else '' for h in header_raw]
print(f"Header: {header}")

def col_idx(names):
    """Return index of first matching column name (case-insensitive)"""
    for name in names:
        for i, h in enumerate(header):
            if h.lower() == name.lower():
                return i
    return -1

idx_projekt  = col_idx(['Projekt', 'projekt'])
idx_dmc      = col_idx(['DMC', 'dmc', 'Dmc'])
idx_typ      = col_idx(['Typ opravy', 'TyptOpravy', 'TypOpravy', 'Typ_opravy', 'Typ opravy'])
# Fallback: if typ not found by name, use index 2 (3rd column per known Excel structure)
if idx_typ < 0:
    idx_typ = 2
    print(f"  NOTE: 'Typ opravy' column not found by name, using positional index 2 ({header[2] if len(header) > 2 else '?'})")
idx_akce     = col_idx(['Akce', 'akce'])
idx_datum    = col_idx(['Datum', 'datum'])
idx_pozn     = col_idx(['Poznámka/Nové QR', 'Poznamka/NoveQR', 'Poznámka', 'poznamka', 'Poznamka'])
idx_nove_qr  = col_idx(['Nové QR', 'NoveQR', 'Nove QR'])

print(f"\nColumn mapping:")
print(f"  Projekt: {idx_projekt} | DMC: {idx_dmc} | Typ opravy: {idx_typ} | Akce: {idx_akce}")
print(f"  Datum: {idx_datum} | Poznámka: {idx_pozn} | Nové QR: {idx_nove_qr}")

if idx_dmc < 0:
    print("ERROR: DMC column not found")
    sys.exit(1)

# ── Process rows ───────────────────────────────────────────────────────────────
def map_project(val):
    if val is None: return None
    s = str(val).strip().lower()
    if s in projects_by_code: return projects_by_code[s]
    # Partial match
    for key, pid in projects_by_code.items():
        if s in key or key in s:
            return pid
    return None

def map_reason(val):
    if val is None: return None
    s = str(val).strip()
    s_low = s.lower()
    if s_low in reasons_by_code: return reasons_by_code[s_low]
    # Try leading number
    num_match = re.match(r'^(\d+)', s)
    if num_match:
        num = num_match.group(1)
        if num in reasons_by_code: return reasons_by_code[num]
        num_stripped = num.lstrip('0') or '0'
        if num_stripped in reasons_by_code: return reasons_by_code[num_stripped]
    # Partial label match
    for key, rid in reasons_by_code.items():
        if len(key) > 2 and key in s_low:
            return rid
    return None

def map_action(val):
    if val is None: return None
    s = str(val).strip().lower()
    if s in actions_by_code: return actions_by_code[s]
    for key, aid in actions_by_code.items():
        if s in key or key in s:
            return aid
    return None

insert_sql = '''
    INSERT INTO rework_entries
        (project_id, dmc, reason_id, action_id, replacement_dmc, note,
         recorded_at, recorded_by_user_id, recorded_by_name)
    VALUES (%s, %s, %s, %s, %s, %s, %s, NULL, %s)
'''

ok_count   = 0
skip_count = 0
warn_count = 0
skip_reasons = {}

BATCH_SIZE = 500
batch = []

def flush_batch():
    global ok_count
    if not batch or DRY_RUN:
        ok_count += len(batch)
        batch.clear()
        return
    try:
        cursor.executemany(insert_sql, batch)
        conn.commit()
        ok_count += len(batch)
    except mysql.connector.Error as e:
        print(f"  DB error flushing batch: {e}")
        conn.rollback()
    batch.clear()

print(f"\nProcessing rows {'(DRY RUN)' if DRY_RUN else ''}...")

for row_num, row in enumerate(rows_iter, start=2):
    if all(v is None for v in row):
        continue

    dmc = safe(row[idx_dmc] if idx_dmc >= 0 else None)
    if not dmc:
        skip_count += 1
        skip_reasons['empty_dmc'] = skip_reasons.get('empty_dmc', 0) + 1
        continue

    projekt_raw = safe(row[idx_projekt] if idx_projekt >= 0 else None)
    typ_raw     = safe(row[idx_typ]     if idx_typ >= 0 else None)
    akce_raw    = safe(row[idx_akce]    if idx_akce >= 0 else None)
    datum_raw   = row[idx_datum] if idx_datum >= 0 else None
    pozn_raw    = safe(row[idx_pozn]    if idx_pozn >= 0 else None)
    nove_qr_raw = safe(row[idx_nove_qr] if idx_nove_qr >= 0 else None)

    project_id = map_project(projekt_raw)
    reason_id  = map_reason(typ_raw)
    action_id  = map_action(akce_raw)
    datum      = safe_date(datum_raw) or datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    if project_id is None:
        # Default to first project if only one or cannot map
        project_id = list(projects_by_code.values())[0] if projects_by_code else None
        warn_count += 1

    if reason_id is None:
        skip_count += 1
        skip_reasons['no_reason'] = skip_reasons.get('no_reason', 0) + 1
        if skip_count <= 5:
            print(f"  Row {row_num}: skip - no reason match for '{typ_raw}'")
        continue

    if action_id is None:
        skip_count += 1
        skip_reasons['no_action'] = skip_reasons.get('no_action', 0) + 1
        if skip_count <= 5:
            print(f"  Row {row_num}: skip - no action match for '{akce_raw}'")
        continue

    if project_id is None:
        skip_count += 1
        skip_reasons['no_project'] = skip_reasons.get('no_project', 0) + 1
        continue

    # Split replacement_dmc / note from Poznámka/Nové QR column if separate col absent
    replacement_dmc = nove_qr_raw
    note            = pozn_raw
    if replacement_dmc is None and note is not None:
        # If note looks like a DMC (alphanumeric, 8-30 chars, no spaces), treat as replacement
        if re.match(r'^[A-Za-z0-9_\-]{8,30}$', note):
            replacement_dmc = note
            note = None

    batch.append((
        project_id, dmc, reason_id, action_id,
        replacement_dmc, note,
        datum,
        'import_rework_excel.py',
    ))

    if len(batch) >= BATCH_SIZE:
        flush_batch()
        print(f"  ...{ok_count} rows imported so far")

flush_batch()
wb.close()

print(f"\n{'='*50}")
print(f"{'DRY RUN - no data written' if DRY_RUN else 'Import complete'}")
print(f"  Imported : {ok_count}")
print(f"  Skipped  : {skip_count}")
print(f"  Warnings : {warn_count} (project defaulted)")
if skip_reasons:
    print(f"  Skip reasons: {skip_reasons}")
print(f"{'='*50}")

cursor.close()
conn.close()
